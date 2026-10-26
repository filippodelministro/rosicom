"""Conversione degli export Excel di Leghe Fantacalcio in squadre.csv e partite.csv.

Entrambi i fogli affiancano più blocchi sulla stessa riga:
- rose: due squadre per riga (colonne A-D e F-I), "Ruolo | Calciatore | Squadra | Costo"
- calendario: due giornate per riga (colonne A-E e G-K), "casa | pt casa | pt trasf | trasf | 2-1"
"""
import csv
import re
from pathlib import Path

RUOLI = {"P", "D", "C", "A"}
RE_GIORNATA = re.compile(r"(\d+)ª Giornata lega")
RE_RISULTATO = re.compile(r"^(\d+)-(\d+)$")


def _cella(riga: list[str], i: int) -> str:
    return riga[i].strip() if i < len(riga) else ""


def griglia_da_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    def testo(v):
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    foglio = load_workbook(path, read_only=True).worksheets[0]
    return [[testo(v) for v in riga] for riga in foglio.iter_rows(values_only=True)]


def rose(griglia: list[list[str]]) -> list[dict]:
    giocatori = []
    for inizio in (0, 5):
        squadra = None
        for i, riga in enumerate(griglia):
            primo = _cella(riga, inizio)
            succ = griglia[i + 1] if i + 1 < len(griglia) else []
            if primo and _cella(succ, inizio) == "Ruolo":
                squadra = primo
            elif primo.startswith("Crediti"):
                squadra = None
            elif squadra and primo in RUOLI:
                giocatori.append({
                    "squadra": squadra, "giocatore": _cella(riga, inizio + 1), "ruolo": primo,
                    "club": _cella(riga, inizio + 2), "costo": _cella(riga, inizio + 3),
                })
    return giocatori


def calendario(griglia: list[list[str]]) -> list[dict]:
    partite = []
    for inizio in (0, 6):
        giornata = None
        for riga in griglia:
            primo = _cella(riga, inizio)
            m = RE_GIORNATA.match(primo)
            if m:
                giornata = int(m.group(1))
            elif giornata and primo:
                ris = RE_RISULTATO.match(_cella(riga, inizio + 4))
                partite.append({
                    "giornata": giornata, "casa": primo, "trasferta": _cella(riga, inizio + 3),
                    "punti_casa": _cella(riga, inizio + 1) if ris else "",
                    "punti_trasferta": _cella(riga, inizio + 2) if ris else "",
                    "gol_casa": ris.group(1) if ris else "",
                    "gol_trasferta": ris.group(2) if ris else "",
                })
    return sorted(partite, key=lambda p: p["giornata"])


def scrivi_csv(righe: list[dict], path: Path) -> None:
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(righe[0]))
        w.writeheader()
        w.writerows(righe)
